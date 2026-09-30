"""
Auto-generated notice/complaint letter generator for KnowYourRights.
HARD PRIVACY CONSTRAINT:
This module executes strictly via local application template substitution.
User personal identifying fields (names, addresses, amounts) are NEVER
sent to any LLM or external provider.
"""

import io
from datetime import date
from typing import Optional, Dict, Any, List
from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.letter_templates import CLAIM_SENTENCE_TEMPLATES, TENANCY_CATEGORIES


class LetterFields(BaseModel):
    employee_name: Optional[str] = Field(default=None, description="Full legal name of employee or tenant")
    employee_address: Optional[str] = Field(default=None, description="Residential or postal address")
    counterparty_name: Optional[str] = Field(default=None, description="Name of employer, company, or landlord")
    counterparty_address: Optional[str] = Field(default=None, description="Address of employer or landlord")
    joining_date: Optional[str] = Field(default=None, description="Joining date or lease commencement date")
    designation: Optional[str] = Field(default=None, description="Job title or role")
    amount_claimed: Optional[str] = Field(default=None, description="Unpaid salary, deposit amount, or claim value")


def is_letter_eligible(
    status: str,
    has_direct_recourse: bool,
    category: Optional[str],
    citations: List[Any],
    evidence_level: Optional[str] = None
) -> bool:
    """
    A case is eligible for an auto-generated formal notice letter ONLY if:
    1. Status is fully resolved (no clarification pending, no error).
    2. hasDirectRecourse is True (the provisions directly cover the scenario).
    3. Evidence level is NOT Low.
    4. Category has a verified statutory claim template.
    5. At least one verified statutory citation exists to cite.
    """
    if status != "resolved":
        return False
    if has_direct_recourse is not True:
        return False
    if evidence_level and evidence_level.strip().lower() == "low":
        return False
    if not category or category not in CLAIM_SENTENCE_TEMPLATES:
        return False
    if not citations or len(citations) == 0:
        return False
    return True


def _build_letter_context(
    category: str,
    act: str,
    section: str,
    authority: str,
    ask_text: str,
    fields: LetterFields,
    today_str: Optional[str] = None
) -> Dict[str, str]:
    current_date = today_str or date.today().strftime("%d %B %Y")

    emp_name = (fields.employee_name or "").strip() or "[Full Legal Name]"
    emp_addr = (fields.employee_address or "").strip() or "[Your Complete Postal Address, City, State, PIN]"
    cp_name = (fields.counterparty_name or "").strip() or "[Employer / Company / Landlord Full Legal Name]"
    cp_addr = (fields.counterparty_address or "").strip() or "[Registered Office Address / Landlord Postal Address, PIN]"

    designation_str = (fields.designation or "").strip() or "[Your Designation / Job Title / Department]"
    period_or_amount = (fields.amount_claimed or "").strip() or "[Claim Amount in INR, e.g., Rs. ________ / Duration of Non-Payment]"

    template = CLAIM_SENTENCE_TEMPLATES.get(
        category,
        "My statutory rights under Section {section} of the {act} have been violated."
    )
    claim_sentence = template.format(
        act=act or "Applicable Indian Statute",
        section=section or "Statutory Section",
        period_or_amount=period_or_amount,
        unpaid_period=period_or_amount,
        amount_claimed=period_or_amount
    )

    clean_ask = (ask_text or "").strip()
    if clean_ask.endswith("."):
        clean_ask = clean_ask[:-1]
    if clean_ask:
        clean_ask = clean_ask[0].lower() + clean_ask[1:]
    else:
        clean_ask = "the outstanding statutory dues and grievances"

    subject_category = category.replace("_", " ").title()
    is_tenancy = category in TENANCY_CATEGORIES

    if is_tenancy:
        joining_str = (fields.joining_date or "").strip() or "[Lease Commencement Date]"
        relationship_sentence = (
            f"I, {emp_name}, am / was a tenant residing at the leased premises administered by you since {joining_str}."
        )
    else:
        joining_str = (fields.joining_date or "").strip() or "[Joining / Commencement Date]"
        relationship_sentence = (
            f"I, {emp_name}, was employed with your organization as {designation_str} since {joining_str}."
        )

    return {
        "current_date": current_date,
        "emp_name": emp_name,
        "emp_addr": emp_addr,
        "cp_name": cp_name,
        "cp_addr": cp_addr,
        "relationship_sentence": relationship_sentence,
        "claim_sentence": claim_sentence,
        "clean_ask": clean_ask,
        "subject_category": subject_category,
        "act": act,
        "section": section,
        "authority": authority,
    }


def render_letter(
    category: str,
    act: str,
    section: str,
    authority: str,
    ask_text: str,
    fields: LetterFields,
    today_str: Optional[str] = None
) -> Dict[str, str]:
    """
    Renders both a plain text notice letter and a printable HTML document.
    Unprovided values have clean, standard legal bracketed placeholders.
    """
    ctx = _build_letter_context(category, act, section, authority, ask_text, fields, today_str)

    text_content = f"""Date: {ctx['current_date']}

To,
{ctx['cp_name']}
{ctx['cp_addr']}

Subject: Formal Notice regarding {ctx['subject_category']} under Section {ctx['section']}, {ctx['act']}

Dear Sir / Madam,

{ctx['relationship_sentence']}

{ctx['claim_sentence']}

I hereby request that {ctx['clean_ask']} be resolved within 15 days of receipt of this notice, failing which I will approach {ctx['authority']} for statutory enforcement, compensation, interest, and formal legal redressal.

Please treat this as formal written notice under the law.

Yours sincerely,

{ctx['emp_name']}
{ctx['emp_addr']}
[Contact Number & Email ID]
"""

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <title>Formal Statutory Notice — {ctx['subject_category']}</title>
  <style>
    @page {{
      size: A4 portrait;
      margin: 25mm 20mm;
    }}
    body {{
      font-family: 'Times New Roman', Times, Georgia, serif;
      font-size: 12pt;
      line-height: 1.6;
      color: #111;
      margin: 0;
      padding: 24px;
      background: #fff;
    }}
    .header-date {{
      text-align: right;
      font-weight: bold;
      margin-bottom: 24px;
    }}
    .recipient-block {{
      margin-bottom: 24px;
      white-space: pre-line;
    }}
    .subject-line {{
      font-weight: bold;
      text-decoration: underline;
      margin-bottom: 20px;
    }}
    .salutation {{
      margin-bottom: 16px;
    }}
    .letter-body {{
      text-align: justify;
      margin-bottom: 16px;
    }}
    .closing {{
      margin-top: 36px;
    }}
    .sender-block {{
      margin-top: 24px;
      white-space: pre-line;
    }}
    .statute-badge {{
      border-top: 1px solid #999;
      margin-top: 40px;
      padding-top: 8px;
      font-family: Arial, sans-serif;
      font-size: 8pt;
      color: #555;
      text-transform: uppercase;
      letter-spacing: 0.05em;
    }}
  </style>
</head>
<body>
  <div class="header-date">{ctx['current_date']}</div>

  <div class="recipient-block">
    <strong>To,</strong><br />
    {ctx['cp_name']}<br />
    {ctx['cp_addr']}
  </div>

  <div class="subject-line">
    Subject: Formal Statutory Notice regarding {ctx['subject_category']} under Section {ctx['section']} of {ctx['act']}
  </div>

  <div class="salutation">Dear Sir / Madam,</div>

  <p class="letter-body">
    {ctx['relationship_sentence']}
  </p>

  <p class="letter-body">
    <strong>{ctx['claim_sentence']}</strong>
  </p>

  <p class="letter-body">
    I hereby formally request that {ctx['clean_ask']} be completely settled and resolved within fifteen (15) days from the receipt of this notice, failing which I shall be constrained to initiate formal proceedings before the <strong>{ctx['authority']}</strong> for enforcement of statutory rights, damages, interest, and costs without further reference to you.
  </p>

  <p class="letter-body">
    Please treat this communication as formal written statutory notice under the law.
  </p>

  <div class="closing">
    Yours sincerely,
    <div class="sender-block">
      <strong>{ctx['emp_name']}</strong><br />
      {ctx['emp_addr']}<br />
      [Contact Number & Email ID]
    </div>
  </div>

  <div class="statute-badge">
    KnowYourRights Statutory Notice Draft · Reference: {ctx['act']} §{ctx['section']} · Legal Recourse Forum: {ctx['authority']}
  </div>
</body>
</html>"""

    return {
        "text": text_content.strip(),
        "html": html_content.strip()
    }


def render_letter_docx(
    category: str,
    act: str,
    section: str,
    authority: str,
    ask_text: str,
    fields: LetterFields,
    today_str: Optional[str] = None
) -> io.BytesIO:
    """
    Renders a formatted Microsoft Word (.docx) document with proper margins,
    typography, and legal notice structure.
    """
    ctx = _build_letter_context(category, act, section, authority, ask_text, fields, today_str)

    doc = Document()

    # 1 inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Base font setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    # Date (right-aligned)
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_after = Pt(16)
    r_date = p_date.add_run(ctx['current_date'])
    r_date.bold = True

    # Recipient
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(14)
    r_to_label = p_to.add_run("To,\n")
    r_to_label.bold = True
    p_to.add_run(f"{ctx['cp_name']}\n{ctx['cp_addr']}")

    # Subject line
    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_after = Pt(14)
    r_subj = p_subj.add_run(f"Subject: Formal Statutory Notice regarding {ctx['subject_category']} under Section {ctx['section']} of {ctx['act']}")
    r_subj.bold = True
    r_subj.underline = True

    # Salutation
    p_sal = doc.add_paragraph("Dear Sir / Madam,")
    p_sal.paragraph_format.space_after = Pt(12)

    # Paragraph 1: Relationship
    p_rel = doc.add_paragraph(ctx['relationship_sentence'])
    p_rel.paragraph_format.space_after = Pt(12)
    p_rel.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Paragraph 2: Claim
    p_claim = doc.add_paragraph()
    p_claim.paragraph_format.space_after = Pt(12)
    p_claim.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_claim = p_claim.add_run(ctx['claim_sentence'])
    r_claim.bold = True

    # Paragraph 3: Demand & 15-day notice
    p_demand = doc.add_paragraph(
        f"I hereby formally request that {ctx['clean_ask']} be completely settled and resolved within fifteen (15) days from the receipt of this notice, failing which I shall be constrained to initiate formal proceedings before the {ctx['authority']} for enforcement of statutory rights, compensation, interest, and costs without further reference to you."
    )
    p_demand.paragraph_format.space_after = Pt(12)
    p_demand.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Paragraph 4: Formal treatment
    p_treat = doc.add_paragraph("Please treat this communication as formal written statutory notice under the law.")
    p_treat.paragraph_format.space_after = Pt(24)

    # Closing & Signature
    p_close = doc.add_paragraph("Yours sincerely,")
    p_close.paragraph_format.space_after = Pt(28)

    p_sig = doc.add_paragraph()
    r_name = p_sig.add_run(f"{ctx['emp_name']}\n")
    r_name.bold = True
    p_sig.add_run(f"{ctx['emp_addr']}\n[Contact Number & Email ID]")
    p_sig.paragraph_format.space_after = Pt(36)

    # Footer note
    p_foot = doc.add_paragraph(f"KnowYourRights Statutory Notice Draft · Reference: {ctx['act']} §{ctx['section']} · Legal Recourse Forum: {ctx['authority']}")
    p_foot.paragraph_format.space_before = Pt(24)
    if p_foot.runs:
        p_foot.runs[0].font.size = Pt(8.5)
        p_foot.runs[0].font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
